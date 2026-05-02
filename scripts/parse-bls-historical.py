#!/usr/bin/env python3
"""
Parse BLS OEWS annual XLSX bundles (2020-2023) into salary.db wages table.

The schema already supports multi-year data via PRIMARY KEY (soc_code, area_code, year).
This script APPENDS years 2020-2023 alongside the existing 2024 data.

Idempotent: deletes wages rows for years 2020-2023 before re-inserting.

Filtering:
  - I_GROUP == 'cross-industry' (NAICS 000000)
  - O_GROUP == 'detailed' (specific 6-digit SOC codes — matches occupations table)
  - Areas matched against existing areas table (50 MSAs + 1 national)

Area code normalization:
  XLSX MSA AREA = '35620'  → DB area_code '0035620' (zero-pad to 7 chars)
  XLSX nat AREA = '99'     → DB area_code '0000000'

Usage: python3 scripts/parse-bls-historical.py
"""
import os
import sqlite3
import sys
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
DB_PATH = ROOT / 'data' / 'salary.db'
HIST_DIR = ROOT / 'data' / 'bls-historical'
YEARS = [2020, 2021, 2022, 2023]


def parse_value(v):
    """BLS uses '*' for suppressed, '#' for top-coded ($208k+/yr or $100+/hr)."""
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip()
    if s in ('', '*', '**', '#', '##'):
        return None
    try:
        return float(s.replace(',', ''))
    except ValueError:
        return None


def normalize_area(area_raw, area_type):
    """Match XLSX area code to DB format.

    XLSX:                          DB (existing):
      AREA='35620', TYPE='4'  →     '0035620'  (MSA)
      AREA='99',    TYPE='1'  →     '0000000'  (national)
    """
    s = str(area_raw).strip()
    if str(area_type).strip() == '1' or s == '99':
        return '0000000'
    # MSA: zero-pad to 7 chars (existing DB uses 7-digit codes)
    return s.zfill(7)


def to_int(v):
    f = parse_value(v)
    return int(round(f)) if f is not None else None


def parse_year(db, year):
    yy = str(year)[-2:]
    msa_path = HIST_DIR / f'oesm{yy}ma' / f'MSA_M{year}_dl.xlsx'
    nat_path = HIST_DIR / f'oesm{yy}nat' / f'national_M{year}_dl.xlsx'

    # Build allowlist sets from existing DB so we only insert rows that match
    # the 50 MSAs + national + 397 occupations our pages render for.
    cur = db.cursor()
    valid_areas = {row[0] for row in cur.execute('SELECT area_code FROM areas')}
    valid_socs = {row[0] for row in cur.execute('SELECT soc_code FROM occupations')}
    print(f'  [{year}] allowlist: {len(valid_areas)} areas, {len(valid_socs)} SOC codes')

    rows_inserted = 0
    rows_skipped_area = 0
    rows_skipped_soc = 0
    rows_skipped_filter = 0

    insert_sql = """
        INSERT OR REPLACE INTO wages
        (soc_code, area_code, employment, annual_mean, annual_median,
         annual_p10, annual_p25, annual_p75, annual_p90,
         hourly_mean, hourly_median, year)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
    """

    for label, path in [('national', nat_path), ('MSA', msa_path)]:
        if not path.exists():
            print(f'  [{year}] SKIP {label}: {path} not found')
            continue
        print(f'  [{year}] reading {label}: {path.name}')
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        rows_iter = ws.iter_rows(values_only=True)
        header = list(next(rows_iter))
        # Map column name → index
        col = {h: i for i, h in enumerate(header)}
        needed = ['AREA', 'AREA_TYPE', 'I_GROUP', 'O_GROUP', 'OCC_CODE',
                  'TOT_EMP', 'A_MEAN', 'A_MEDIAN', 'A_PCT10', 'A_PCT25',
                  'A_PCT75', 'A_PCT90', 'H_MEAN', 'H_MEDIAN']
        for n in needed:
            if n not in col:
                print(f'  ERROR: column {n} missing in {path.name}; have {header}')
                wb.close()
                return 0

        batch = []
        for row in rows_iter:
            if not row or row[col['OCC_CODE']] is None:
                continue
            i_group = (row[col['I_GROUP']] or '').strip().lower()
            o_group = (row[col['O_GROUP']] or '').strip().lower()
            if i_group != 'cross-industry':
                rows_skipped_filter += 1
                continue
            if o_group != 'detailed':
                rows_skipped_filter += 1
                continue

            area_code = normalize_area(row[col['AREA']], row[col['AREA_TYPE']])
            soc_code = (row[col['OCC_CODE']] or '').strip()

            if soc_code not in valid_socs:
                rows_skipped_soc += 1
                continue
            if area_code not in valid_areas:
                rows_skipped_area += 1
                continue

            a_median = to_int(row[col['A_MEDIAN']])
            a_mean = to_int(row[col['A_MEAN']])
            if a_median is None and a_mean is None:
                continue

            batch.append((
                soc_code, area_code,
                to_int(row[col['TOT_EMP']]),
                a_mean, a_median,
                to_int(row[col['A_PCT10']]),
                to_int(row[col['A_PCT25']]),
                to_int(row[col['A_PCT75']]),
                to_int(row[col['A_PCT90']]),
                parse_value(row[col['H_MEAN']]),
                parse_value(row[col['H_MEDIAN']]),
                year,
            ))

            if len(batch) >= 5000:
                cur.executemany(insert_sql, batch)
                rows_inserted += len(batch)
                batch.clear()

        if batch:
            cur.executemany(insert_sql, batch)
            rows_inserted += len(batch)
        wb.close()

    db.commit()
    print(f'  [{year}] inserted {rows_inserted} rows '
          f'(skip area={rows_skipped_area}, soc={rows_skipped_soc}, filter={rows_skipped_filter})')
    return rows_inserted


def main():
    if not DB_PATH.exists():
        print(f'ERROR: {DB_PATH} does not exist. Run parse-bls.ts first.')
        sys.exit(1)

    db = sqlite3.connect(DB_PATH)
    db.execute('PRAGMA journal_mode = WAL')
    db.execute('PRAGMA synchronous = OFF')

    # Idempotent: clear historical years before re-inserting
    db.execute(f'DELETE FROM wages WHERE year IN ({",".join(str(y) for y in YEARS)})')
    db.commit()

    total = 0
    for year in YEARS:
        total += parse_year(db, year)

    # Add year index if missing
    db.execute('CREATE INDEX IF NOT EXISTS idx_wages_year ON wages(year)')
    db.commit()

    print('\n=== Year coverage ===')
    for row in db.execute('SELECT year, COUNT(*) FROM wages GROUP BY year ORDER BY year'):
        print(f'  {row[0]}: {row[1]:>6,} rows')

    print(f'\n  Total backfilled: {total} rows across {len(YEARS)} years')

    db.close()


if __name__ == '__main__':
    main()
